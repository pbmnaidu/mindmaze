from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "docs" / "NIRIKSHAN_REFERENCE_GUIDE.md"
OUT = ROOT / "docs" / "NIRIKSHAN_REFERENCE_GUIDE.docx"
IMG = ROOT / "docs" / "guide_images"
IMG.mkdir(exist_ok=True)

def make_diagram(path, title, labels):
    im = Image.new("RGB", (1500, 280), "white")
    d = ImageDraw.Draw(im)
    try:
        f1 = ImageFont.truetype(r"C:\Windows\Fonts\segoeui.ttf", 20)
        f2 = ImageFont.truetype(r"C:\Windows\Fonts\segoeuib.ttf", 24)
    except OSError:
        f1 = f2 = ImageFont.load_default()
    d.text((30, 20), title, fill="#111827", font=f2)
    width, gap, y = 220, 55, 95
    for i, label in enumerate(labels):
        x = 30 + i * (width + gap)
        d.rounded_rectangle((x, y, x + width, y + 120), radius=12, fill="#EAF2F8", outline="#1F4E79", width=3)
        words, lines, line = label.split(), [], ""
        for word in words:
            if len(line) + len(word) > 18:
                lines.append(line); line = word
            else:
                line = (line + " " + word).strip()
        lines.append(line)
        d.multiline_text((x + 14, y + 28), "\n".join(lines), fill="#111827", font=f1, spacing=5)
        if i < len(labels) - 1:
            ax = x + width + 8
            d.line((ax, y + 60, ax + gap - 18, y + 60), fill="#1F4E79", width=4)
            d.polygon([(ax + gap - 18, y + 60), (ax + gap - 32, y + 50), (ax + gap - 32, y + 70)], fill="#1F4E79")
    im.save(path)

make_diagram(IMG / "pipeline.png", "NIRIKSHAN data to review", ["Source data", "Clean records", "Classify work", "Calculate scores", "Human review"])
make_diagram(IMG / "financial.png", "Financial anomaly path", ["Description", "Quantity", "Expected cost", "Compare", "Reason"])
make_diagram(IMG / "review.png", "Reviewer workflow", ["Open monitor", "Start at top", "Open work", "Check records", "Decide"])

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def add_table(doc, lines):
    rows = [[x.strip() for x in line.strip().strip("|").split("|")] for line in lines if line.strip() and "---" not in line]
    if len(rows) < 2: return
    t = doc.add_table(rows=1, cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(rows[0]):
        t.rows[0].cells[i].text = value; shade(t.rows[0].cells[i], "1F4E79")
        for r in t.rows[0].cells[i].paragraphs[0].runs: r.font.bold = True; r.font.color.rgb = RGBColor(255,255,255); r.font.size = Pt(8)
    for row in rows[1:]:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for r in cells[i].paragraphs[0].runs: r.font.size = Pt(8)
    doc.add_paragraph()

def add_text(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p

doc = Document(); sec = doc.sections[0]; sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.7); sec.right_margin = Inches(.7)
doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(10)
for name, size in [("Title", 25), ("Heading 1", 17), ("Heading 2", 13)]:
    doc.styles[name].font.name = "Aptos Display"; doc.styles[name].font.size = Pt(size); doc.styles[name].font.color.rgb = RGBColor(0,0,0)

lines = MD.read_text(encoding="utf-8").splitlines(); i = 0; in_code = False; table_lines = []
while i < len(lines):
    line = lines[i]
    if line.startswith("```mermaid"):
        while i < len(lines) and lines[i].strip() != "```": i += 1
        i += 1; continue
    if line.startswith("```"):
        in_code = not in_code
        if in_code: add_text(doc, "")
        i += 1; continue
    if in_code:
        p = doc.add_paragraph(); r = p.add_run(line); r.font.name = "Consolas"; r.font.size = Pt(9); r.font.color.rgb = RGBColor(31,78,121); i += 1; continue
    if line.startswith("|"):
        table_lines.append(line); i += 1
        if i >= len(lines) or not lines[i].startswith("|"):
            add_table(doc, table_lines); table_lines = []
        continue
    if line.startswith("# "):
        p = doc.add_paragraph(style="Title"); p.add_run(line[2:])
    elif line.startswith("## "):
        doc.add_heading(line[3:], level=1)
    elif line.startswith("### "):
        doc.add_heading(line[4:], level=2)
    elif line.startswith("- "):
        add_text(doc, line[2:], "List Bullet")
    elif line[:2].isdigit() and line[2:4] == ". ":
        add_text(doc, line[4:], "List Number")
    elif line.strip() and not line.startswith("```"):
        add_text(doc, line)
    if line.startswith("## 2. Complete Workflow"):
        doc.add_picture(str(IMG / "pipeline.png"), width=Inches(6.8)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if line.startswith("## 5. Financial"):
        doc.add_picture(str(IMG / "financial.png"), width=Inches(6.8)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if line.startswith("## 10. Reviewer"):
        doc.add_picture(str(IMG / "review.png"), width=Inches(6.8)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    i += 1

doc.save(OUT); print(OUT)
